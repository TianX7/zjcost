import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "旧建材市场调研表"

thin = Side(style="thin", color="000000")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(bold=True, color="FFFFFF", size=11)
title_font = Font(bold=True, size=18)
subtitle_font = Font(size=10, color="666666")
bold11 = Font(bold=True, size=11)
normal10 = Font(size=10)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_wrap = Alignment(vertical="top", wrap_text=True)

# ── 标题 ──
ws.merge_cells("A1:L1")
ws["A1"] = "旧建材市场调研表"
ws["A1"].font = title_font
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

ws.merge_cells("A2:L2")
ws["A2"] = "编制单位：________________    调研人员：________________    调研日期：________________    调研地点：________________"
ws["A2"].font = subtitle_font
ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[2].height = 22

# ── 表头 ──
headers = [
    ("序号", 6),
    ("建材名称", 16),
    ("规格型号", 16),
    ("材质/品牌", 14),
    ("单位", 8),
    ("单价（元）", 12),
    ("新旧程度", 12),
    ("来源渠道", 14),
    ("供应商/摊位号", 16),
    ("联系方式", 14),
    ("可供应量", 12),
    ("质量评价/备注", 20),
]
for col, (name, width) in enumerate(headers, 1):
    c = ws.cell(row=4, column=col, value=name)
    c.font = header_font
    c.fill = header_fill
    c.alignment = center
    c.border = border
    ws.column_dimensions[get_column_letter(col)].width = width
ws.row_dimensions[4].height = 28

# ── 预填旧建材品类（引导填写） ──
categories = [
    "旧门窗（木门/铝合金门/塑钢门）",
    "旧铝合金型材",
    "旧瓷砖/地砖/墙砖",
    "旧水泥制品（预制板/砌块）",
    "旧钢材（钢管/钢筋/角钢）",
    "旧模板/木方",
    "旧管道（PPR/PVC/铸铁管）",
    "旧电线电缆",
    "旧灯具/开关面板",
    "旧卫浴洁具",
    "旧涂料/油漆（未开封）",
    "旧石材（大理石/花岗岩）",
    "旧玻璃（钢化/中空）",
    "旧保温材料",
    "旧配电箱/开关柜",
    "其他旧建材：",
    "",
    "",
    "",
    "",
]

for i, cat in enumerate(categories):
    row = 5 + i
    ws.cell(row=row, column=1, value=i + 1 if cat else i + 1).alignment = center
    if cat:
        ws.cell(row=row, column=2, value=cat)
    for col in range(1, 13):
        ws.cell(row=row, column=col).border = border
        ws.cell(row=row, column=col).font = normal10
    ws.row_dimensions[row].height = 22

# ── 分隔行 ──
sep_row = 25
ws.merge_cells(f"A{sep_row}:L{sep_row}")
ws.cell(row=sep_row, column=1).border = border
ws.row_dimensions[sep_row].height = 6

# ── 市场信息 ──
info_start = 26
ws.merge_cells(f"A{info_start}:L{info_start}")
ws.cell(row=info_start, column=1, value="市场基本信息").font = Font(bold=True, size=14)
ws.cell(row=info_start, column=1).alignment = Alignment(horizontal="center")
ws.row_dimensions[info_start].height = 30

info_items = [
    ("市场名称", ""),
    ("市场地址", ""),
    ("经营商户数", ""),
    ("主要经营品类", ""),
    ("市场整体新旧程度", "□ 全新  □ 九成新  □ 八成新  □ 七成新及以下"),
    ("价格水平评价", "□ 偏高  □ 适中  □ 偏低"),
    ("物流配套", "□ 自有物流  □ 第三方物流  □ 无物流"),
    ("质量保障", "□ 有质保  □ 部分有质保  □ 无质保"),
    ("综合评价", ""),
    ("调研人签字", "________________    日期：________________"),
    ("审核人签字", "________________    日期：________________"),
]

for i, (label, value) in enumerate(info_items):
    r = info_start + 1 + i
    ws.merge_cells(f"A{r}:B{r}")
    ws.cell(row=r, column=1, value=label).font = bold11
    ws.cell(row=r, column=1).alignment = left_wrap
    ws.cell(row=r, column=1).border = border
    ws.cell(row=r, column=2).border = border
    ws.merge_cells(f"C{r}:L{r}")
    ws.cell(row=r, column=3, value=value).font = normal10
    ws.cell(row=r, column=3).alignment = left_wrap
    ws.cell(row=r, column=3).border = border
    ws.row_dimensions[r].height = 24

path = r"C:\Users\TIAN\Desktop\AI\zjcost-main\旧建材市场调研表.xlsx"
wb.save(path)
print(f"OK: {path}")
